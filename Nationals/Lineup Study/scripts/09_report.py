import json
import re
import subprocess
import time
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "out"
TAB = OUT / "tables"
NAVY, RED, GREY = "#14346B", "#AB0003", "#8A8F98"
SHAPES = {"A": "PPCCPCCPC", "B": "PCPCPCCPC", "C": "CPCPCCPCP"}
ERAS = ["Pre DH 1926 to 1972", "AL DH 1973 to 2021", "NL non DH 1973 to 2021", "Universal DH 2022 to 2026"]
RECALC = Path("/root/.claude/skills/synced/ad1eab3c-92fc-4b58-a820-0eb12a50d243_a877aaed-fea2-4a65-ae0c-016ae19c1431/xlsx/scripts/recalc.py")

plt.rcParams.update({"font.size": 24, "axes.titlesize": 34, "axes.labelsize": 28, "xtick.labelsize": 24,
                     "ytick.labelsize": 24, "legend.fontsize": 22, "axes.spines.top": False, "axes.spines.right": False,
                     "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white"})


def log(msg):
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S} 09_report {msg}"
    print(line)
    with open(ROOT / "logs" / "run.log", "a") as f:
        f.write(line + "\n")


def t(name):
    return pd.read_csv(TAB / f"{name}.csv")


def fig():
    return plt.figure(figsize=(19.2, 10.8), dpi=100)


def foot(f, text):
    f.text(0.01, 0.015, text, fontsize=16, color=GREY)


def chart_a(b):
    f = fig()
    ax = f.add_axes([0.09, 0.14, 0.62, 0.72])
    order = ["0 to 1", "2", "3", "4", "5+"]
    styles = [(ERAS[0], NAVY, "o", "-"), (ERAS[1], "#3F6DB5", "s", "-"), (ERAS[2], "#7FA3D9", "^", "-"),
              (ERAS[3], GREY, "D", "-"), ("2026", RED, "o", "-")]
    for g, c, m, ls in styles:
        x = b[b.group == g].set_index("bucket").reindex(order)
        lw = 6 if g == "2026" else 4
        ax.plot(range(5), x.rot, color=c, marker=m, ms=16, lw=lw, ls=ls, label=g)
    ax.axhline(0, color="black", lw=1.5)
    ax.set_xticks(range(5), ["0 to 1", "2", "3", "4", "5 or more"])
    ax.set_xlabel("Slots that differ from shape A (P P C C P C C P C)")
    ax.set_ylabel("Runs over talent per game")
    ax.legend(loc="center left", bbox_to_anchor=(1.02, 0.5), frameon=False)
    f.suptitle("Lineups closest to shape A scored more only in 2026", x=0.09, ha="left", fontsize=36, fontweight="bold")
    foot(f, "Runs over talent = actual runs minus a within season model of runs on lineup talent, opposing starter, opponent runs allowed, home and park.")
    f.savefig(OUT / "chart_a_runs_over_talent_by_distance.png")
    plt.close(f)


def chart_b(o, s26):
    d = o[(o.test == "oos_decade") & (o["shape"] == "A")].copy()
    comb = o[(o.test == "oos_pooled") & (o["shape"] == "A")].iloc[0]
    f = fig()
    ax = f.add_axes([0.09, 0.17, 0.88, 0.69])
    x = np.arange(len(d))
    v = -d.slope * 162
    se162 = d.se * 162
    ax.errorbar(x, v, yerr=1.96 * se162, fmt="none", ecolor=NAVY, elinewidth=2, alpha=0.5)
    ax.errorbar(x, v, yerr=se162, fmt="o", color=NAVY, ms=18, elinewidth=7)
    xr = len(d) + 0.7
    v26 = -s26.slope * 162
    ax.errorbar([xr], [v26], yerr=[1.96 * s26.se * 162], fmt="none", ecolor=RED, elinewidth=2, alpha=0.5)
    ax.errorbar([xr], [v26], yerr=[s26.se * 162], fmt="o", color=RED, ms=22, elinewidth=7)
    lo, hi = -(comb.slope + 1.96 * comb.se) * 162, -(comb.slope - 1.96 * comb.se) * 162
    ax.axhspan(lo, hi, color=NAVY, alpha=0.12)
    ax.axhline(0, color="black", lw=1.5)
    labels = [g.replace("2020s", "2020\nto 2025") for g in d.group] + ["2026"]
    ax.set_xticks(list(x) + [xr], labels)
    ax.set_ylim(-12, 28)
    ax.set_ylabel("Runs gained per 162 games\nper slot closer to A")
    ax.text(xr - 0.25, v26, "the season\nA was\npicked in", color=RED, ha="right", va="center", fontsize=22)
    ax.text(-0.3, 22, f"1926 to 2025 pooled: {abs(comb.slope * 162):.2f} runs per 162 per slot (shaded band: 95% interval)", color=NAVY, fontsize=22)
    f.suptitle("Out of sample, shape A is worth nothing in every decade", x=0.09, ha="left", fontsize=36, fontweight="bold")
    foot(f, "Above zero = lineups closer to shape A scored more than talent predicts. Thick bars: plus or minus 1 SE; thin bars: 95% interval. SE clustered by team season.")
    f.savefig(OUT / "chart_b_shapeA_slope_by_decade.png")
    plt.close(f)


def chart_c(tr):
    tr = tr.copy()
    tr.loc[tr.lineups * (1 - tr.x_share) < 500, "p_share"] = np.nan
    p = tr.pivot(index="decade", columns="slot", values="p_share") * 100
    f = fig()
    ax = f.add_axes([0.12, 0.15, 0.80, 0.71])
    cmap = LinearSegmentedColormap.from_list("nav", ["#FFFFFF", NAVY])
    cmap.set_bad("#EEEEEE")
    ax.imshow(np.ma.masked_invalid(p.values), cmap=cmap, vmin=0, vmax=100, aspect="auto")
    for i in range(p.shape[0]):
        for j in range(p.shape[1]):
            v = p.values[i, j]
            ax.text(j, i, "pitcher" if np.isnan(v) else f"{v:.0f}", ha="center", va="center", fontsize=22 if not np.isnan(v) else 18,
                    color=GREY if np.isnan(v) else ("white" if v > 55 else "black"))
    for j, c in enumerate(SHAPES["A"]):
        if c == "P":
            ax.add_patch(plt.Rectangle((j - 0.5, -0.5), 1, p.shape[0], fill=False, ec=RED, lw=5))
    ax.set_xticks(range(9), [str(k) for k in range(1, 10)])
    ax.set_yticks(range(p.shape[0]), [d.replace("2020s", "2020 to 26") for d in p.index])
    ax.set_xlabel("Batting order slot (red outline = P slot in shape A)")
    ax.spines[:].set_visible(False)
    f.suptitle("Share of power lean hitters (P) by slot, percent", x=0.12, ha="left", fontsize=36, fontweight="bold")
    foot(f, "P = SLG percentile above contact percentile within that season. Pitchers batting left out; grey = fewer than 500 non pitchers in that slot.")
    f.savefig(OUT / "chart_c_power_share_heatmap.png")
    plt.close(f)


def chart_d(mk, fe):
    groups = ["2026", "1926 to 2025 sample"]
    labels = ["2026", "1926 to 2025"]
    emp = [mk.set_index("group").loc[g] for g in groups]
    f = fig()
    ax = f.add_axes([0.10, 0.14, 0.86, 0.70])
    x = np.arange(2)
    w = 0.36
    ev = np.array([-e.empirical_slope * 162 for e in emp])
    es = np.array([e.empirical_se * 162 for e in emp])
    mv = np.array([e.markov_per_slot_toward_A_calibrated * 162 for e in emp])
    ax.bar(x - w / 2, ev, w, color=NAVY, label="Observed: runs over talent per slot closer to A")
    ax.errorbar(x - w / 2, ev, yerr=1.96 * es, fmt="none", ecolor="black", elinewidth=3, capsize=10)
    ax.bar(x + w / 2, mv, w, color=RED, label="Markov model: pure batting order effect per slot moved to A")
    for xi, v, e in zip(x - w / 2, ev, es):
        ax.text(xi + 0.03, v + 1.96 * e + 0.4 if v >= 0 else 1.96 * e + 0.4, f"{abs(v):.1f}" if abs(v) >= 0.05 else "0.0", ha="left", va="bottom", fontsize=28, color=NAVY, fontweight="bold")
    for xi, v in zip(x + w / 2, mv):
        ax.text(xi, v - 0.5, f"{abs(v):.2f} {'lost' if v < 0 else 'gained'}", ha="center", va="top", fontsize=26, color=RED, fontweight="bold")
    ax.axhline(0, color="black", lw=1.5)
    ax.set_ylim(-4, 27)
    ax.set_xticks(x, labels)
    ax.set_ylabel("Runs gained per 162 games, per slot")
    ax.legend(loc="upper right", frameon=False)
    f.suptitle("Batting order alone is worth a fraction of a run per slot", x=0.10, ha="left", fontsize=36, fontweight="bold")
    foot(f, "Observed: 95% interval, SE clustered by team season; history uses all 1926 to 2025 lineups. Markov: all 2026 lineups and a 50,000 lineup sample, scaled to actual run levels.")
    f.savefig(OUT / "chart_d_empirical_vs_markov.png")
    plt.close(f)


def chart_e(null, perm):
    f = fig()
    ax = f.add_axes([0.09, 0.15, 0.88, 0.69])
    ax.hist(-null.best_slope * 162, bins=40, color=NAVY, alpha=0.85)
    rb, ra = -perm["real_best_slope"] * 162, -perm["real_slope_A"] * 162
    ax.axvline(rb, color=RED, lw=6)
    ax.axvline(ra, color=RED, lw=4, ls="--")
    top = ax.get_ylim()[1] * 1.3
    ax.set_ylim(0, top)
    box = dict(facecolor="white", edgecolor="none", alpha=0.9)
    ax.text(rb + 0.3, top * 0.98, f"real 2026 best of 512\n{rb:.1f} runs per 162", color=RED, fontsize=22, va="top", ha="left", bbox=box)
    ax.text(ra - 0.3, top * 0.98, f"shape A\n{ra:.1f} runs per 162", color=RED, fontsize=22, va="top", ha="right", bbox=box)
    ax.set_xlabel("Best shape in each shuffle: runs gained per 162 per slot closer to it")
    ax.set_ylabel("Shuffles")
    pv = perm["p_value_best"]
    f.suptitle(f"Shuffled tags find a 2026 winner this strong {pv:.0%} of the time", x=0.09, ha="left", fontsize=36, fontweight="bold")
    foot(f, f"{perm['n_perm']:,} shuffles of P and C tags among each 2026 team's own hitters; the best of 512 shapes is kept from each shuffle.")
    f.savefig(OUT / "chart_e_permutation_null.png")
    plt.close(f)


HFILL = PatternFill("solid", fgColor="14346B")
HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BFONT = Font(name="Arial", size=10)


def write_table(ws, df, start_col=1, formulas=None, formats=None, filt=True):
    formulas = formulas or {}
    formats = formats or {}
    cols = list(df.columns) + list(formulas)
    for j, c in enumerate(cols):
        cell = ws.cell(row=1, column=start_col + j, value=c)
        cell.fill, cell.font = HFILL, HFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    letter = {c: get_column_letter(start_col + j) for j, c in enumerate(cols)}
    vals = df.to_numpy(dtype=object)
    for i in range(len(df)):
        r = i + 2
        for j, v in enumerate(vals[i]):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            elif isinstance(v, (np.bool_, bool)):
                v = bool(v)
            cell = ws.cell(row=r, column=start_col + j, value=v)
            cell.font = BFONT
            if df.columns[j] in formats:
                cell.number_format = formats[df.columns[j]]
        for k, (c, ftmp) in enumerate(formulas.items()):
            f = ftmp.format(r=r, **letter)
            cell = ws.cell(row=r, column=start_col + len(df.columns) + k, value=f)
            cell.font = BFONT
            if c in formats:
                cell.number_format = formats[c]
    for j, c in enumerate(cols):
        ws.column_dimensions[get_column_letter(start_col + j)].width = max(10, min(42, len(str(c)) + 4))
    if filt:
        ws.auto_filter.ref = f"{get_column_letter(start_col)}1:{get_column_letter(start_col + len(cols) - 1)}{len(df) + 1}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42
    return letter


def workbook(H, tables):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = tables["summary"]
    write_table(ws, s, formulas={"Runs per 162": "=IF(ISNUMBER({Runs per game}{r}),{Runs per game}{r}*162,\"\")",
                                 "SE runs per 162": "=IF(ISNUMBER({SE runs per game}{r}),{SE runs per game}{r}*162,\"\")",
                                 "Wins per 162": "=IF(ISNUMBER({Runs per game}{r}),{Runs per game}{r}*162/{Runs per win}{r},\"\")"},
                formats={"Runs per game": "0.0000", "SE runs per game": "0.0000", "Runs per 162": "0.0", "SE runs per 162": "0.0",
                         "Wins per 162": "0.00", "Runs per win": "0.00"})
    r = len(s) + 3
    ws.cell(row=r, column=1, value="Verdict").font = Font(name="Arial", bold=True, size=11, color="14346B")
    ws.cell(row=r, column=2, value=H["verdict"]).font = Font(name="Arial", size=11)
    ws.cell(row=r + 1, column=1, value="Benchmark").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=r + 1, column=2, value="Published research puts the whole gap between the best and a typical batting order at 5 to 15 runs a season.").font = BFONT
    ws.cell(row=r + 2, column=1, value="Sign").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=r + 2, column=2, value="Negative slope = lineups closer to the shape scored more runs than their talent predicted.").font = BFONT
    ws.column_dimensions["B"].width = 60

    ws = wb.create_sheet("Shape Tests")
    sl = tables["slopes"]
    L = write_table(ws, sl, formulas={"t stat": "={slope_runs_per_game}{r}/{se}{r}",
                                      "Runs per 162 (formula)": "={slope_runs_per_game}{r}*162",
                                      "Wins per 162 (formula)": "={slope_runs_per_game}{r}*162/{runs_per_win}{r}"},
                    formats={"slope_runs_per_game": "0.0000", "se": "0.0000", "t stat": "0.00", "runs_per_win": "0.00",
                             "Runs per 162 (formula)": "0.0", "Wins per 162 (formula)": "0.00", "mean_dist": "0.00"})
    b = tables["buckets"]
    write_table(ws, b, start_col=len(sl.columns) + 5,
                formulas={"Win pct": "={wins}{r}/{decisions}{r}", "R per G": "={runs}{r}/{games}{r}",
                          "Runs over talent per G": "=({runs}{r}-{exp_runs}{r})/{games}{r}"},
                formats={"Win pct": "0.000", "R per G": "0.00", "Runs over talent per G": "0.000", "exp_runs": "0.0"}, filt=False)

    ws = wb.create_sheet("Order vs Composition")
    write_table(ws, tables["markov"], formulas={"Shape A minus actual per 162 (formula)": "={markov_shapeA}{r}*162-{markov_actual}{r}*162",
                                                 "Best minus actual per 162 (formula)": "={markov_best}{r}*162-{markov_actual}{r}*162",
                                                 "Observed runs per 162 per slot (formula)": "=-{empirical_slope}{r}*162",
                                                 "Markov runs per 162 per slot (formula)": "={markov_per_slot_toward_A_calibrated}{r}*162",
                                                 "Observed over Markov (formula)": "=IF({markov_per_slot_toward_A_calibrated}{r}<>0,-{empirical_slope}{r}/{markov_per_slot_toward_A_calibrated}{r},\"\")"},
                formats={c: "0.0000" for c in tables["markov"].columns if c not in ("group", "lineups")} | {
                    "Shape A minus actual per 162 (formula)": "0.00", "Best minus actual per 162 (formula)": "0.00",
                    "Observed runs per 162 per slot (formula)": "0.00", "Markov runs per 162 per slot (formula)": "0.00",
                    "Observed over Markov (formula)": "0.0"})

    ws = wb.create_sheet("512 Shapes")
    write_table(ws, tables["shapes"], formulas={"t stat": "={slope}{r}/{se}{r}", "Runs per 162": "={slope}{r}*162"},
                formats={"slope": "0.0000", "se": "0.0000", "mean_rot_within_2": "0.000", "t stat": "0.00", "Runs per 162": "0.0"})

    ws = wb.create_sheet("Exact Shapes")
    write_table(ws, tables["exact"], formulas={"Win pct": "=IF({W}{r}+{L}{r}>0,{W}{r}/({W}{r}+{L}{r}),\"\")",
                                                "R per G": "={R}{r}/{G}{r}", "RA per G": "={RA}{r}/{G}{r}",
                                                "Runs over talent per G": "=({R}{r}-{exp_R}{r})/{G}{r}"},
                formats={"exp_R": "0.0", "Win pct": "0.000", "R per G": "0.00", "RA per G": "0.00", "Runs over talent per G": "0.000"})

    ws = wb.create_sheet("By Decade")
    tr = tables["trends"]
    write_table(ws, tr, formats={"p_share": "0.0%", "x_share": "0.0%", "mean_slg": "0.000", "mean_k_pct": "0.0%", "mean_obp": "0.000"})
    write_table(ws, tables["dist"], start_col=len(tr.columns) + 3, filt=False,
                formulas={"Distance to A per 9 counted slots (formula)": "={mean_dist_A}{r}/{mean_counted}{r}*9"},
                formats={"mean_dist_A": "0.00", "mean_counted": "0.00", "mean_dist_B": "0.00", "mean_dist_C": "0.00",
                         "mean_dist_A_per_9": "0.00", "Distance to A per 9 counted slots (formula)": "0.00"})

    ws = wb.create_sheet("Teams")
    write_table(ws, tables["teams"], formulas={"Pythag W (formula)": "=({W}{r}+{L}{r})*{R}{r}^1.83/({R}{r}^1.83+{RA}{r}^1.83)",
                                                "W minus Pythag (formula)": "={W}{r}-({W}{r}+{L}{r})*{R}{r}^1.83/({R}{r}^1.83+{RA}{r}^1.83)",
                                                "Win pct": "={W}{r}/({W}{r}+{L}{r})"},
                formats={"mean_dist_A": "0.00", "rot_g": "0.000", "pyth_w": "0.0", "w_minus_pyth": "0.0",
                         "Pythag W (formula)": "0.0", "W minus Pythag (formula)": "0.0", "Win pct": "0.000"})

    ws = wb.create_sheet("Examples")
    write_table(ws, tables["examples"], formulas={"Win pct": "={W}{r}/({W}{r}+{L}{r})"}, formats={"Win pct": "0.000"})
    for c in range(1, 30):
        ws.column_dimensions[get_column_letter(c)].width = 16
    for k in range(1, 10):
        col = list(tables["examples"].columns).index(f"slot{k}") + 1
        ws.column_dimensions[get_column_letter(col)].width = 44

    ws = wb.create_sheet("Validation and Method")
    write_table(ws, tables["method"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path = OUT / "Lineup_Study_1926_2026.xlsx"
    wb.save(path)
    return path


def recalc(path):
    if RECALC.exists():
        r = subprocess.run(["python3", str(RECALC), str(path), "300"], capture_output=True, text=True)
        try:
            return json.loads(r.stdout.strip().splitlines()[-1] if r.stdout.strip().startswith("{") is False else r.stdout)
        except Exception:
            return {"error": r.stdout[-500:] + r.stderr[-500:]}
    tmp = OUT / "recalc_tmp"
    tmp.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp), str(path)], capture_output=True)
    (tmp / path.name).replace(path)
    tmp.rmdir()
    return {"status": "converted with soffice"}


def formula_errors(path):
    wb = load_workbook(path, data_only=True)
    bad, n = 0, 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad += 1
                n += 1
    return bad


def words(v, pos="more", neg="fewer", fmt="{:.1f}"):
    return f"{fmt.format(abs(v))} {pos if v >= 0 else neg}"


def cluster_ols(x, y, g):
    m = sm.OLS(np.asarray(y, float), sm.add_constant(np.asarray(x, float))).fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(g)[0]})
    return m.params[1], m.bse[1]


def second_way(H, R, ex, mk, null):
    lu = pd.read_parquet(DATA / "lineup_model.parquet", columns=["season", "pattern", "dist_A", "rot", "team_season", "win", "runs",
                                                                  "runs_allowed", "dh_game", "game_key", "home"])
    nx, nn = lu.pattern.str.count("X"), lu.pattern.str.count("N")
    lu = lu[((nx == 0) | ((nx == 1) & ~lu.dh_game)) & (nn == 0)]
    d26 = lu[lu.season == 2026]
    rows = []
    b, se = cluster_ols(d26.dist_A, d26.rot, d26.team_season)
    rows.append(("2026 shape A slope (runs per game per slot)", H["slope_A_2026"], b, "512 shape aggregate formula vs statsmodels OLS on lineups"))
    rows.append(("2026 shape A SE", H["se_A_2026"], se, "aggregate sandwich vs statsmodels cluster SE"))
    h = lu[lu.season <= 2025]
    x = h.dist_A.to_numpy(float) - h.dist_A.mean()
    rows.append(("1926 to 2025 shape A slope", R["oos_pooled_A"]["slope"], float((x * h.rot.to_numpy()).sum() / (x * x).sum()), "statsmodels vs numpy covariance over variance"))
    lineups = pd.read_parquet(DATA / "lineups.parquet")
    ps = pd.read_parquet(DATA / "player_seasons.parquet", columns=["season", "player_id", "tag"])
    tag = ps[ps.season == 2026].set_index("player_id").tag
    l26 = lineups[(lineups.season == 2026) & lineups.complete]
    pat = l26[[f"p{k}_id" for k in range(1, 10)]].apply(lambda c: c.map(tag)).fillna("N").agg("".join, axis=1)
    a = l26[pat == SHAPES["A"]]
    e26 = ex[(ex.group == "2026") & (ex.pattern == SHAPES["A"])].iloc[0]
    rows.append(("2026 exact shape A games", e26.G, len(a), "exact shape table vs tags rebuilt from player seasons"))
    rows.append(("2026 exact shape A wins", e26.W, int((a.runs > a.runs_allowed).sum()), "exact shape table vs scores in lineups"))
    rows.append(("Permutation p value for the best of 512", R["permutation"]["p_value_best"], float((null.best_slope <= R["permutation"]["real_best_slope"]).mean()), "json vs recount of the null file"))
    m = pd.read_parquet(DATA / "markov_lineups.parquet", columns=["season", "m_actual", "m_shapeA", "m_best", "slots_moved_toward_A", "runs"])
    m26 = m[m.season == 2026]
    mv = m26[m26.slots_moved_toward_A > 0]
    dA = (mv.m_shapeA - mv.m_actual).to_numpy()
    k = mv.slots_moved_toward_A.to_numpy(float)
    reg = float((dA * k).sum() / (k * k).sum())
    rows.append(("2026 Markov per slot toward A (uncalibrated)", mk.set_index("group").loc["2026", "markov_per_slot_toward_A"], reg, "ratio of sums vs regression through the origin (close, not identical by design)"))
    rows.append(("2026 Markov best minus actual per 162", mk.set_index("group").loc["2026", "best_minus_actual_per_162"], float((m26.m_best - m26.m_actual).mean() * 162), "summary table vs markov file"))
    g = pd.read_parquet(DATA / "games.parquet", columns=["season", "vis_score", "home_score"])
    g26 = g[g.season == 2026]
    rpw_games = 10 * np.sqrt((g26.vis_score + g26.home_score).mean() / 9)
    s = t("shape_slopes")
    rows.append(("2026 runs per win", s[(s.group == "2026") & (s["sample"] == "main") & (s["shape"] == "A")].runs_per_win.iloc[0], rpw_games, "lineups table vs games table"))
    y = d26.rot - d26.groupby("team_season").rot.transform("mean")
    xx = d26.dist_A - d26.groupby("team_season").dist_A.transform("mean")
    rows.append(("2026 within team season slope", R["fe_A_2026"]["slope"], float((xx * y).sum() / (xx * xx).sum()), "statsmodels demeaned vs numpy within"))
    c = pd.DataFrame(rows, columns=["number", "first_way", "second_way", "how"])
    c["abs_diff"] = (c.first_way - c.second_way).abs()
    c["ok"] = c.abs_diff <= np.maximum(1e-6, 0.02 * c.first_way.abs())
    c.loc[c.number.str.contains("Markov per slot"), "ok"] = c.abs_diff <= 0.25 * c.first_way.abs()
    return c


def build_tables(H, R):
    sl = t("shape_slopes")
    o = t("out_of_sample")
    fe = t("fixed_effects")
    mk = t("markov_summary")
    wm = t("win_models")
    rb = t("robustness")
    perm = R["permutation"]
    rpw26 = sl[(sl.group == "2026") & (sl["sample"] == "main")].runs_per_win.iloc[0]
    rpwh = o[(o.test == "oos_pooled")].runs_per_win.iloc[0]

    def row(metric, sample, rg=None, se=None, rpw=None, other=None, other_se=None, note="", src=""):
        return {"Metric": metric, "Sample": sample, "Runs per game": rg, "SE runs per game": se, "Runs per win": rpw,
                "Other value": other, "Other SE": other_se, "Note": note, "Source file": src}
    S = []
    for s in "ABC":
        S.append(row(f"Shape {s} slope, runs over talent per slot off", "2026, main sample", H[f"slope_{s}_2026"], H[f"se_{s}_2026"], rpw26,
                     H[f"rank_{s}_2026"], None, f"Other value = rank of 512 (1 = best)", "out/tables/all_512_shapes.csv"))
    for s in "ABC":
        r = o[(o.test == "oos_pooled") & (o["shape"] == s)].iloc[0]
        S.append(row(f"Shape {s} slope, out of sample", "1926 to 2025 pooled", r.slope, r.se, r.runs_per_win, None, None, "", "out/tables/out_of_sample.csv"))
    r = o[(o.test == "oos_combined_inverse_variance") & (o["shape"] == "A")].iloc[0]
    S.append(row("Shape A slope, decades combined", "1926 to 2025 inverse variance", r.slope, r.se, r.runs_per_win, None, None, "", "out/tables/out_of_sample.csv"))
    for g in ["1926 to 2025", "2026"]:
        r = fe[(fe.test == "team_season_fixed_effects") & (fe.group == g) & (fe["shape"] == "A")].iloc[0]
        S.append(row("Shape A slope within team season", g, r.slope, r.se, r.runs_per_win, None, None, "team season fixed effects", "out/tables/fixed_effects.csv"))
    for g in ["2026", "1926 to 2025 sample"]:
        r = mk[mk.group == g].iloc[0]
        S.append(row("Markov pure order effect per slot moved toward A", g, r.markov_per_slot_toward_A_calibrated, None, rpw26 if g == "2026" else rpwh,
                     None, None, "scaled to actual run level", "out/tables/markov_summary.csv"))
        S.append(row("Markov: rearrange into shape A minus actual order", g, r.shapeA_minus_actual_per_game * r.calibration, None, rpw26 if g == "2026" else rpwh,
                     r.share_shapeA_beats_actual, None, "Other value = share of lineups where shape A order is better", "out/tables/markov_summary.csv"))
        S.append(row("Markov: best order minus actual order", g, r.best_minus_actual_per_game * r.calibration, None, rpw26 if g == "2026" else rpwh,
                     None, None, "local search over swaps", "out/tables/markov_summary.csv"))
    for g in ["2026", "1926 to 2026"]:
        r = wm[(wm.group == g) & (wm["shape"] == "A")].iloc[0]
        S.append(row("Win model, wins per 162 per slot off A (linear)", g, None, None, None, r.lpm_wins_per_162, r.lpm_se_wins_per_162, "linear probability model", "out/tables/win_models.csv"))
        S.append(row("Win model, wins per 162 per slot off A (logistic)", g, None, None, None, r.logit_ame_wins_per_162, r.logit_ame_se_wins_per_162, "average marginal effect", "out/tables/win_models.csv"))
    S.append(row("Permutation: share of shuffles with a best shape at least as strong", "2026", None, None, None, perm["p_value_best"], None,
                 f"{perm['n_perm']} shuffles within team season", "out/tables/permutation_null.csv"))
    S.append(row("Shape A t statistic vs Bonferroni bar for 512 shapes", "2026", None, None, None, H["slope_A_2026"] / H["se_A_2026"], perm["bonferroni_z"],
                 "Other SE column holds the bar (z for 0.05 over 512 tests)", "out/tables/headline.json"))
    S.append(row("Exact shape A lineups, wins and losses", "2026", None, None, None, H["exact_A_2026_W"], H["exact_A_2026_L"],
                 f"{H['exact_A_2026_G']} games; Other value = W, Other SE = L", "out/tables/exact_shapes.csv"))
    S.append(row("Shape A slope with xwOBA talent", "2026", R["slope_A_2026_xwoba"], rb[(rb.test == "2026 xwOBA talent") & (rb["shape"] == "A")].se.iloc[0], rpw26,
                 R["rank_A_2026_xwoba"], None, "Other value = rank of 512", "out/tables/shapes_2026_xwoba.csv"))
    summary = pd.DataFrame(S)
    shapes = t("all_512_shapes")
    ex = t("exact_shapes")
    ex["flag"] = ""
    for gname, x in ex[ex.G >= 100].groupby("group"):
        top = x.sort_values("win_pct", ascending=False)
        ex.loc[top.head(10).index, "flag"] = "Top 10 at 100+ games"
        ex.loc[top.tail(5).index, "flag"] = "Bottom 5 at 100+ games"
    ex = ex.drop(columns=["win_pct", "r_g", "ra_g", "rot_g"])
    teams = t("team_seasons").drop(columns=["team_season"])
    examples = t("examples")
    v = (OUT / "VALIDATION.md").read_text()
    method = [("Data: 1926 to 2025 lineups and scores", "Retrosheet game logs from the Chadwick Bureau GitHub mirror (retrosheet.org was blocked)."),
              ("Data: 1926 to 2025 player seasons", "Lahman database 2026 edition (through 2025) from the Lahman R package repository."),
              ("Data: 2026", "Stats API live feed pitch mirror (github.com/lancebroz/mlb-pitcher-data) and box score mirror (github.com/toolstem/daily-box-score). Lineups rebuilt from plate appearance order."),
              ("Validation: team seasons", re.search(r"Exact match on all four[^\n]*", v).group(0)),
              ("Validation: 2026 records", re.search(r"Wins and losses vs the final 2026 standings[^\n]*", v).group(0)),
              ("Validation: spot check", re.search(r"Games matching on date, score and all nine slots[^\n]*", v).group(0)),
              ("Validation: completeness", re.search(r"Team games with nine identified starters[^\n]*", v).group(0)),
              ("Tags", "P when the shrunk SLG percentile beats the contact percentile (100 minus shrunk K% percentile), within the season's 250+ PA pool. Priors: 200 PA for SLG, 60 PA for K%. Pitchers batting = X."),
              ("Shapes", "A = P P C C P C C P C; B = P C P C P C C P C; C = C P C P C C P C P. Distance = slots that differ, skipping X slots."),
              ("Talent", "Slot weighted mean of shrunk relative OPS (OBP/lgOBP + SLG/lgSLG minus 1), weights 4.65 to 3.75."),
              ("Runs over talent", "Actual runs minus fitted runs from a within season OLS on talent, opponent starter FIP, opponent RA per game, home and park dummies."),
              ("Standard errors", "Clustered by team season."),
              ("Runs per win", "10 times the square root of (runs per game by both teams / 9) for the group."),
              ("Markov model", "24 base out states, batter carried across innings, 9 innings, event rates 1B, 2B, 3B, HR, BB+HBP, K, other out with a 100 PA prior; runner on 2nd scores on a single 60%, runner on 1st to 3rd 30%, runner on 1st scores on a double 40%, double play on 22% of other outs with a runner on 1st and fewer than 2 outs, productive outs 30%."),
              ("Benchmark", "The Book (Tango, Lichtman, Dolphin), Ruane, Bukiet, Harold and Palacios (1997), Sokol (2003): best vs typical order worth 5 to 15 runs a season."),
              ("Full detail", "VALIDATION.md, DECISIONS.md, PROBLEMS.md and README.md in the study folder.")]
    method = pd.DataFrame(method, columns=["Item", "Detail"])
    return {"summary": summary, "slopes": sl, "buckets": t("distance_buckets"), "markov": mk, "shapes": shapes, "exact": ex,
            "trends": t("trends_by_slot_decade"), "dist": t("distance_by_decade"), "teams": teams, "examples": examples, "method": method}


def report_text(H, R, T):
    sl, o, fe, mk, wm, card = t("shape_slopes"), t("out_of_sample"), t("fixed_effects"), t("markov_summary"), t("win_models"), t("manager_card")
    perm = R["permutation"]
    BT = json.loads((ROOT / "raw" / "manual" / "brief_targets.json").read_text())
    rpw26 = sl[(sl.group == "2026") & (sl["sample"] == "main") & (sl["shape"] == "A")].runs_per_win.iloc[0]
    a26, se26 = H["slope_A_2026"] * 162, H["se_A_2026"] * 162
    oo = o[(o.test == "oos_pooled") & (o["shape"] == "A")].iloc[0]
    oa, ose = oo.slope * 162, oo.se * 162
    lo, hi = oa - 1.96 * ose, oa + 1.96 * ose
    dec = o[(o.test == "oos_decade") & (o["shape"] == "A")]
    max_t = dec.t.abs().max()
    fh = fe[(fe.test == "team_season_fixed_effects") & (fe.group == "1926 to 2025") & (fe["shape"] == "A")].iloc[0]
    m26 = mk[mk.group == "2026"].iloc[0]
    mh = mk[mk.group == "1926 to 2025 sample"].iloc[0]
    per_slot26 = m26.markov_per_slot_toward_A_calibrated * 162
    ratio = abs(a26 / per_slot26)
    w_all = wm[(wm.group == "1926 to 2026") & (wm["shape"] == "A")].iloc[0]
    mean_dist26 = sl[(sl.group == "2026") & (sl["sample"] == "main") & (sl["shape"] == "A")].mean_dist.iloc[0]
    c26 = card[card.group == "2026"].set_index("slot")
    lineups_h = int(oo.lineups)
    A = "P P C C P C C P C"
    verdict = (f"Shape A ({A}) does not hold up: its 2026 edge is the kind of result that picking the best of 512 shapes in one season produces by chance, "
               f"across {lineups_h:,} lineups from 1926 to 2025 it is worth {abs(oa):.2f} runs per 162 games per slot, in effect zero, and batting order on its own is worth well under one run per slot.")
    H["verdict"] = verdict
    lines = [
        "# Power vs contact batting order, 1926 to 2026", "",
        f"**Verdict: {verdict}**", "",
        "## Is the shape real?", "",
        f"* In 2026 each slot closer to shape A came with {abs(a26):.1f} more runs per 162 games than lineup talent predicted (SE {se26:.1f}). That is the pattern the shape was picked from. In this rebuild A ranked {H['rank_A_2026']} of 512 shapes, B {H['rank_B_2026']} and C {H['rank_C_2026']}.",
        f"* Picking the best of 512 shapes inflates the winner. When P and C tags are shuffled among each team's own hitters, the best shuffled shape is at least as strong as the real 2026 best {perm['p_value_best']:.0%} of the time ({perm['n_perm']:,} shuffles). Shape A's t statistic of {abs(H['slope_A_2026'] / H['se_A_2026']):.1f} is below the {perm['bonferroni_z']:.1f} bar that 512 tries require.",
        f"* Every season from 1926 to 2025 is out of sample. There the effect is {abs(oa):.2f} runs per 162 per slot (SE {ose:.2f}), and the 95% interval runs from {abs(lo):.1f} runs {'gained' if lo < 0 else 'lost'} to {abs(hi):.1f} runs {'lost' if hi > 0 else 'gained'} for each slot closer to A. No decade gets past {max_t:.1f} SE, and comparing a team only with its own other games gives {abs(fh.slope * 162):.2f} runs per 162 per slot ({'against' if fh.slope > 0 else 'for'} A).",
        "",
        "## How big in wins?", "",
        f"* Taken at face value, 2026 says {abs(a26) / rpw26:.2f} wins per 162 per slot ({rpw26:.1f} runs per win). The average 2026 lineup sat {mean_dist26:.1f} slots from A, so full adoption would be worth about {abs(a26) * mean_dist26:.0f} runs. The published research puts the gap between the best and a typical order at 5 to 15 runs for the whole order, so the 2026 number is several times too large to be a batting order effect.",
        f"* History says {abs(oa) / oo.runs_per_win:.3f} wins per 162 per slot. The win model over all 1926 to 2026 lineups gives {abs(w_all.lpm_wins_per_162):.2f} wins per 162 per slot (SE {w_all.lpm_se_wins_per_162:.2f}).",
        f"* The Markov run model, which holds the nine hitters fixed and only changes the order, says the best possible order beats the order managers actually used by {m26.best_minus_actual_per_162 * m26.calibration:.1f} runs per 162 in 2026 and {mh.best_minus_actual_per_162 * mh.calibration:.1f} runs from 1926 to 2025, inside the 5 to 15 run benchmark. Rearranging the same hitters into shape A {'costs' if m26.shapeA_minus_actual_per_game < 0 else 'adds'} {abs(m26.shapeA_minus_actual_per_game * m26.calibration * 162):.1f} runs per 162 in 2026. Per slot moved toward A the pure order effect is {abs(per_slot26):.2f} runs per 162 ({'fewer' if per_slot26 < 0 else 'more'}), so the 2026 slope is about {ratio:.0f} times larger and points the other way. The 2026 result comes from which hitters were in the lineup and from chance, not from their order.",
        "",
        "## What a manager should do", "",
        "* Do not set lineups by a power and contact pattern. No shape, A, B or C, moves runs by more than about one run per 162 per slot in 100 years of games.",
        f"* Order by on base skill first, then power. In the Markov best orders for 2026, slots 2, 3 and 4 hold power lean hitters {c26.loc[2, 'p_share_markov_best']:.0%}, {c26.loc[3, 'p_share_markov_best']:.0%} and {c26.loc[4, 'p_share_markov_best']:.0%} of the time, while slots 1, 8 and 9 lean contact (P share {c26.loc[1, 'p_share_markov_best']:.0%}, {c26.loc[8, 'p_share_markov_best']:.0%}, {c26.loc[9, 'p_share_markov_best']:.0%}). The leadoff and number 2 hitters carry the highest OBP ({c26.loc[1, 'obp_markov_best']:.3f} and {c26.loc[2, 'obp_markov_best']:.3f}).",
        f"* Do not sweat the bottom. Swapping neighbors in slots 6 to 9 costs {c26.loc[[6, 7, 8], 'swap_with_next_cost_per_162'].min():.1f} to {c26.loc[[6, 7, 8], 'swap_with_next_cost_per_162'].max():.1f} runs per 162; even the costliest swap at the top costs {c26.swap_with_next_cost_per_162.max():.1f} runs.",
        "",
        "## Caveats", "",
        f"* The MLB Stats API, Savant and Baseball Reference were blocked where this ran. 2026 lineups, scores and batting lines were rebuilt from public mirrors of the Stats API feed; wins and losses match the final standings for all 30 teams. The rebuild gives shape A {abs(H['slope_A_2026']):.3f} runs per game per slot (rank {H['rank_A_2026']}) against {abs(BT['slope_A']):.3f} (rank {BT['rank_A']}) in the earlier analysis quoted in the brief, and the same {H['exact_A_2026_G']} exact shape A games at {H['exact_A_2026_W']} and {H['exact_A_2026_L']}. With xwOBA talent A ranks {R['rank_A_2026_xwoba']}. The conclusion does not depend on these small differences.",
        f"* The Markov model leaves out steals, errors and extra innings, so it reads {(m26.calibration - 1):.0%} low in 2026 and {(mh.calibration - 1):.0%} low historically; its order effects are scaled up by those factors.",
        "* Tags are relative to each season's hitters, and lineup talent is measured with the same season's stats, so a hitter's tag and talent share the noise of that season.",
        "",
        "Every number above is computed by `scripts/09_report.py` from files in `out/tables/` and checked a second way in `out/CHECKS.md`.",
    ]
    return "\n".join(lines) + "\n"


def slides_text(H, R):
    o, mk, card = t("out_of_sample"), t("markov_summary"), t("manager_card")
    perm = R["permutation"]
    oo = o[(o.test == "oos_pooled") & (o["shape"] == "A")].iloc[0]
    m26 = mk[mk.group == "2026"].iloc[0]
    c26 = card[card.group == "2026"].set_index("slot")
    ch = card[card.group != "2026"].set_index("slot")
    cost = c26.swap_with_next_cost_per_162
    firm_cut = cost.median()
    rows = []
    for k in range(1, 10):
        p = c26.loc[k, "p_share_markov_best"]
        lean = "power" if p >= 0.6 else ("contact" if p <= 0.4 else "either")
        near = [cost.get(k - 1, np.nan), cost.get(k, np.nan)]
        near = [v for v in near if not np.isnan(v)]
        status = "firm" if np.mean(near) >= firm_cut else "bends"
        rows.append(f"| {k} | {lean} ({p:.0%} P in best orders; {ch.loc[k, 'p_share_markov_best']:.0%} in 1973 to 2025 DH lineups) | OBP {c26.loc[k, 'obp_markov_best']:.3f} | {status} (neighbor swap {np.mean(near):.2f} runs per 162) |")
    L = ["# Slides", "",
         "## Slide 1: The headline rule", "",
         "**Title:** No batting order shape beats talent, and shape A was a one season pattern", "",
         "**Chart:** `out/chart_b_shapeA_slope_by_decade.png`", "",
         f"* 2026: {abs(H['slope_A_2026']) * 162:.1f} runs per 162 per slot closer to A, the season the shape was picked from.",
         f"* 1926 to 2025: {abs(oo.slope) * 162:.2f} runs per 162 per slot across {int(oo.lineups):,} lineups, flat in every decade.",
         f"* Shuffled tags produce a best shape this strong {perm['p_value_best']:.0%} of the time.", "",
         "**Speaker note:** A was chosen as the best of 512 patterns in one season. The other 100 seasons are a fair test, and the pattern is not there.", "",
         "## Slide 2: Order vs composition", "",
         "**Title:** Batting order alone is worth a fraction of a run per slot", "",
         "**Chart:** `out/chart_d_empirical_vs_markov.png`", "",
         f"* Moving the same nine hitters one slot toward A {'costs' if m26.markov_per_slot_toward_A_calibrated < 0 else 'adds'} {abs(m26.markov_per_slot_toward_A_calibrated * 162):.2f} runs per 162.",
         f"* The best possible order beats what managers used by {m26.best_minus_actual_per_162 * m26.calibration:.1f} runs per 162, at the low end of the 5 to 15 run benchmark.",
         "* The 2026 edge came from who was in the lineup and from chance, not from where they batted.", "",
         "## Slide 3: The manager card", "",
         "**Title:** Best on base hitters first, power in 2 to 4, and do not sweat 6 to 9", "",
         "| Slot | Lean | Typical OBP in best orders | Firm or bends |", "|---|---|---|---|"] + rows + [
         "",
         f"Firm = average cost of swapping with a neighbor at or above the median ({firm_cut:.2f} runs per 162). All swaps are small: the whole best order is worth about {m26.best_minus_actual_per_162 * m26.calibration:.0f} runs a season over current practice.", ""]
    return "\n".join(L)


DASH = re.compile(r"(?<![|\s])[-‐-―−](?![|\s-]*\|)")


def dash_issues(text):
    bad = []
    for i, line in enumerate(text.splitlines(), 1):
        if line.startswith("|"):
            continue
        clean = re.sub(r"`[^`]*`", "", line)
        if re.search(r"[\u2010-\u2015\u2212]", clean) or re.search(r"(?<=\w)-(?=\w)", clean) or re.search(r"(^|\s)-(\s|\d)", clean):
            bad.append((i, line[:100]))
    return bad


def main():
    t0 = time.time()
    H = json.loads((TAB / "headline.json").read_text())
    R = json.loads((TAB / "rigor_headline.json").read_text())
    chart_a(t("distance_buckets"))
    sl = t("all_512_shapes")
    s26 = sl[(sl.group == "2026") & (sl["shape"] == SHAPES["A"])].iloc[0]
    chart_b(t("out_of_sample"), s26)
    chart_c(t("trends_by_slot_decade"))
    chart_d_data = t("markov_summary").copy()
    o = t("out_of_sample")
    full = o[(o.test == "oos_pooled") & (o["shape"] == "A")].iloc[0]
    i = chart_d_data.group == "1926 to 2025 sample"
    chart_d_data.loc[i, "empirical_slope"], chart_d_data.loc[i, "empirical_se"] = full.slope, full.se
    chart_d(chart_d_data, t("fixed_effects"))
    null = t("permutation_null")
    chart_e(null, R["permutation"])
    log("charts written")
    rep = report_text(H, R, None)
    (OUT / "REPORT.md").write_text(rep)
    (OUT / "SLIDES.md").write_text(slides_text(H, R))
    checks = second_way(H, R, t("exact_shapes"), t("markov_summary"), null)
    L = ["# Second way checks", "", "Each headline number recomputed by an independent route.", "",
         "| Number | First way | Second way | Difference | Agrees | How |", "|---|---|---|---|---|---|"]
    for _, r in checks.iterrows():
        L.append(f"| {r.number} | {r.first_way:.6g} | {r.second_way:.6g} | {r.abs_diff:.2g} | {'yes' if r.ok else 'NO'} | {r.how} |")
    (OUT / "CHECKS.md").write_text("\n".join(L) + "\n")
    tables = build_tables(H, R)
    path = workbook(H, tables)
    rc = recalc(path)
    errs = formula_errors(path)
    wb = load_workbook(path)
    empty = [ws.title for ws in wb.worksheets if ws.max_row < 2]
    log(f"workbook recalc {rc.get('status', rc)} total_formulas {rc.get('total_formulas')} errors {rc.get('total_errors')} scan errors {errs} empty tabs {empty}")
    (TAB / "workbook_check.json").write_text(json.dumps({"recalc": rc, "error_cells_scan": errs, "empty_tabs": empty,
                                                         "tabs": [ws.title for ws in wb.worksheets]}, indent=2, default=str))
    issues = {f: dash_issues((OUT / f).read_text()) for f in ["REPORT.md", "SLIDES.md", "VALIDATION.md"]}
    log(f"dash issues: { {k: len(v) for k, v in issues.items()} }; checks ok {int(checks.ok.sum())} of {len(checks)}; in {time.time() - t0:.0f}s")
    (TAB / "dash_check.json").write_text(json.dumps(issues, indent=2))


if __name__ == "__main__":
    main()
